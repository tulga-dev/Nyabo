"""Bytes of an uploaded statement -> rows of cells.

Only the file container is handled here; what the columns mean is the layout's job
(`nyabo_mn.core.statements`). Cells keep their native type (openpyxl gives datetimes
and numbers for typed cells) because the layout parser converts them with the bank's
own date formats instead of guessing.

Legacy `.xls` (BIFF) needs xlrd, which Frappe does not ship; rather than a half-working
reader the accountant gets a clear Mongolian message asking for `.xlsx`.
"""

from __future__ import annotations

import csv
import io
import os
from typing import Any

from nyabo_mn.i18n import mn

XLSX_EXTENSIONS = (".xlsx", ".xlsm")
XLS_EXTENSIONS = (".xls",)
CSV_EXTENSIONS = (".csv", ".txt", ".tsv")
# utf-8-sig first so a BOM never survives as a header character; cp1251 last because it
# decodes almost any byte sequence and would shadow a real UTF-8 file.
CSV_ENCODINGS: tuple[str, ...] = ("utf-8-sig", "utf-8", "cp1251")
CSV_DELIMITERS = ",;\t|"
# Bytes that identify the containers regardless of the extension the bank used.
ZIP_MAGIC = b"PK\x03\x04"
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


class StatementFileError(ValueError):
	"""The upload cannot be read; ``message_mn`` is what the bot shows."""

	def __init__(self, message_mn: str, *, kind: str = "unreadable"):
		super().__init__(message_mn)
		self.message_mn = message_mn
		self.kind = kind


def _extension(filename: str) -> str:
	return os.path.splitext(filename or "")[1].lower()


def read_rows(data: bytes, filename: str) -> list[list[Any]]:
	"""Rows of the first non-empty sheet (xlsx) or of the delimited text file (csv).

	The container is sniffed from the bytes first, so a `.csv` that is really an xlsx
	(banks do that) still opens. Trailing empty cells are kept: column indexes in the
	layout must stay stable across rows.
	"""
	if not data:
		raise StatementFileError(mn.MSG_STATEMENT_FILE_UNREADABLE.format(filename=filename or ""))
	ext = _extension(filename)
	if data.startswith(OLE_MAGIC) or ext in XLS_EXTENSIONS:
		raise StatementFileError(mn.MSG_STATEMENT_FILE_XLS_UNSUPPORTED, kind="xls")
	if data.startswith(ZIP_MAGIC) or ext in XLSX_EXTENSIONS:
		return read_xlsx(data, filename)
	if ext in CSV_EXTENSIONS or not ext:
		return read_csv(data, filename)
	raise StatementFileError(mn.MSG_STATEMENT_FILE_UNREADABLE.format(filename=filename or ""))


def read_xlsx(data: bytes, filename: str = "") -> list[list[Any]]:
	try:
		from openpyxl import load_workbook
	except ImportError as exc:  # pragma: no cover - openpyxl ships with Frappe
		raise StatementFileError(mn.MSG_STATEMENT_FILE_UNREADABLE.format(filename=filename)) from exc
	try:
		workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
	except Exception as exc:  # noqa: BLE001 - openpyxl raises several unrelated types for a bad zip
		raise StatementFileError(mn.MSG_STATEMENT_FILE_UNREADABLE.format(filename=filename)) from exc
	try:
		for sheet in workbook.worksheets:
			rows = [list(row) for row in sheet.iter_rows(values_only=True)]
			rows = _trim_trailing_empty_rows(rows)
			if rows:
				return rows
	finally:
		workbook.close()
	return []


def decode_csv(data: bytes) -> tuple[str, str]:
	"""(text, encoding). utf-8 wins when it decodes; cp1251 is the Windows export fallback."""
	for encoding in CSV_ENCODINGS:
		try:
			return data.decode(encoding), encoding
		except UnicodeDecodeError:
			continue
	return data.decode("cp1251", errors="replace"), "cp1251"


def sniff_delimiter(text: str) -> str:
	"""The delimiter whose per-line count is the same on the most lines.

	``csv.Sniffer`` is not used: a title row such as "Хаан банк, хуулга" makes it pick the
	comma for a semicolon file. Consistency across the data rows is what identifies the
	real separator; ties go to the larger column count.
	"""
	lines = [line for line in text.splitlines()[:50] if line.strip()]
	best, best_key = ",", (0, 0)
	for candidate in CSV_DELIMITERS:
		counts = [line.count(candidate) for line in lines]
		nonzero = [c for c in counts if c > 0]
		if not nonzero:
			continue
		mode = max(set(nonzero), key=nonzero.count)
		key = (nonzero.count(mode), mode)
		if key > best_key:
			best, best_key = candidate, key
	return best


def read_csv(data: bytes, filename: str = "") -> list[list[Any]]:
	text, _encoding = decode_csv(data)
	if not text.strip():
		return []
	delimiter = sniff_delimiter(text)
	reader = csv.reader(io.StringIO(text), delimiter=delimiter)
	rows = [[cell.strip() for cell in row] for row in reader]
	return _trim_trailing_empty_rows(rows)


def _trim_trailing_empty_rows(rows: list[list[Any]]) -> list[list[Any]]:
	while rows and all(cell in (None, "") for cell in rows[-1]):
		rows.pop()
	return rows


def preview_rows(rows: list[list[Any]], limit: int = 5, width: int = 40) -> list[list[str]]:
	"""First rows as short strings for the "unknown layout" card."""
	out: list[list[str]] = []
	for row in rows[:limit]:
		cells = []
		for cell in row:
			text = "" if cell is None else str(cell).strip()
			cells.append(text if len(text) <= width else text[: width - 1] + "…")
		out.append(cells)
	return out


__all__ = [
	"CSV_ENCODINGS",
	"StatementFileError",
	"decode_csv",
	"preview_rows",
	"read_csv",
	"read_rows",
	"read_xlsx",
	"sniff_delimiter",
]
