"""Simulate one receipt under both regimes, or one bank statement, without network or posting.

	bench --site <site> execute nyabo_mn.simulator.run.run --kwargs '{"case": "petrovis_fuel"}'
	python -m nyabo_mn.simulator.run petrovis_fuel
	python -m nyabo_mn.simulator.run --statement path/to/statement.xlsx

``simulate_receipt`` prints the two proposed entries side by side (VAT payer on 2026-06-15,
simplified 1% on 2027-02-15) with the explanation, the citation and every reason the
proposal would be held for the accountant. It sets ``frappe.flags.nyabo_simulation``,
uses ``MockLlmClient`` (fixtures) and the mock ebarimt provider when one exists, and —
when Frappe is importable — provisions a throwaway ``SIM-`` company per regime through
``nyabo_mn.setup.provision_company`` so the chart and the settings rows are the real
ones, then cleans up. Nothing is ever submitted: the simulator stops where the card
would be shown.
"""

from __future__ import annotations

import argparse
import datetime as dt
import importlib
import json
import sys
from collections.abc import Iterable, Sequence
from pathlib import Path
from typing import Any

from nyabo_mn.agent.mock_client import MockLlmClient
from nyabo_mn.core.models import Regime
from nyabo_mn.core.money import fmt_mnt
from nyabo_mn.evals import harness
from nyabo_mn.evals.loader import GOLDEN_DIR, EvalCase, load_file
from nyabo_mn.i18n import mn

# The names come from the Regime enum that ``rules.regime`` re-exports (F-12); the
# simulator runs without a site, so it reads the enum rather than importing the Frappe-side
# module.
REGIME_VAT_PAYER = Regime.VAT_PAYER.value
REGIME_SIMPLIFIED = Regime.SIMPLIFIED_1PCT.value
DEFAULT_REGIMES: tuple[str, ...] = (REGIME_VAT_PAYER, REGIME_SIMPLIFIED)
DEFAULT_DATES: tuple[str, ...] = ("2026-06-15", "2027-02-15")
SIM_PREFIX = "SIM-"
COLUMN_WIDTH = 52
REGIME_LABELS = {
	REGIME_VAT_PAYER: mn.SIM_REGIME_VAT_PAYER,
	REGIME_SIMPLIFIED: mn.SIM_REGIME_SIMPLIFIED,
}


# --- inputs ------------------------------------------------------------------------------------


def load_case(case_or_path: str | Path | EvalCase | dict[str, Any]) -> EvalCase:
	"""A golden case id (``petrovis_fuel``), a JSON file path, a case dict, or an EvalCase."""
	if isinstance(case_or_path, EvalCase):
		return case_or_path
	if isinstance(case_or_path, dict):
		from nyabo_mn.evals.loader import validate_case

		return validate_case(case_or_path, path="inline")
	text = str(case_or_path)
	path = Path(text)
	if path.suffix == ".json" and path.is_file():
		cases = load_file(path)
		if not cases:
			raise ValueError(f"{path}: no cases")
		return cases[0]
	# Prefer the extraction case (the vision call runs), then a classification case.
	by_id: dict[str, EvalCase] = {}
	for file in sorted(GOLDEN_DIR.glob("*.json")):
		for case in load_file(file):
			by_id[case.case_id] = case
	for candidate in (text, f"extract_{text}", f"classify_{text}_vat_payer"):
		if candidate in by_id:
			return by_id[candidate]
	raise LookupError(f"no golden case named {text!r}")


def _receipt_from_case(case: EvalCase, client: MockLlmClient) -> tuple[dict[str, Any], dict[str, Any]]:
	"""(receipt_dict, extra input) — extraction cases go through the vision call first."""
	input_json = dict(case.input_json)
	if input_json.get("receipt"):
		fixture_key = str(input_json.get("classify_fixture") or "")
		for regime in DEFAULT_REGIMES:
			if fixture_key.endswith(f"_{regime}"):
				stem = fixture_key[: -len(regime) - 1]
				input_json.setdefault(
					"classify_fixture_by_regime", {r: f"{stem}_{r}" for r in DEFAULT_REGIMES}
				)
		return dict(input_json["receipt"]), input_json
	fixture = str(input_json.get("llm_fixture") or "extract/default")
	outcome = harness.extract_case(client, fixture)
	receipt = dict(outcome.receipt_dict)
	stem = fixture.rsplit("/", 1)[-1]
	input_json.setdefault("classify_fixture_by_regime", {r: f"classify/{stem}_{r}" for r in DEFAULT_REGIMES})
	input_json.setdefault("seller_vat_payer", True if receipt.get("vat_amount") else None)
	return receipt, input_json


# --- seller lookup ----------------------------------------------------------------------------------


def seller_lookup(receipt: dict[str, Any]) -> tuple[bool | None, str]:
	"""Mock ebarimt registry: ``nyabo_mn.ebarimt.mock.MockProvider`` when present, else the receipt itself."""
	try:
		mock = importlib.import_module("nyabo_mn.ebarimt.mock")
		provider = mock.MockProvider()
		info = provider.lookup_seller(
			tin=receipt.get("seller_tin"), register_no=receipt.get("seller_register_no")
		)
		return getattr(info, "vat_payer", None), str(getattr(provider, "name", "mock"))
	except (ImportError, AttributeError):
		return (True if receipt.get("vat_amount") else None), "receipt"


# --- companies ------------------------------------------------------------------------------------


def _frappe(use_site: bool | None = None) -> Any | None:
	"""The frappe module when a connected site should be used, else None (pure-Python run)."""
	if use_site is False:
		return None
	try:
		import frappe
	except ImportError:
		return None
	connected = getattr(frappe, "local", None) is not None and getattr(frappe.local, "db", None) is not None
	return frappe if connected else None


def provision_sim_company(regime: str, use_site: bool | None = None) -> str | None:
	"""``SIM-<regime>`` through the real provisioning code; None when there is no site.

	``provision_company`` writes the Nyabo Company Settings and the first regime row itself
	(``vat_registered`` decides the regime, effective from the fiscal-year start), so the
	simulator adds nothing on top: a second regime row would only overlap the real one.
	"""
	frappe = _frappe(use_site)
	if frappe is None:
		return None
	from nyabo_mn.setup.provision_company import provision_company

	name = f"{SIM_PREFIX}{regime}"
	abbr = "SV" if regime == REGIME_VAT_PAYER else "SS"
	report = provision_company(name, abbr, vat_registered=1 if regime == REGIME_VAT_PAYER else 0)
	if not report["verify"]["ok"]:
		raise RuntimeError(f"simulation company {name} did not verify: {report['verify']['problems']}")
	return name


def cleanup_sim_company(name: str) -> list[str]:
	"""Delete what the simulation created; refusals are reported, never raised.

	UNVERIFIED: ERPNext's own ``Company.on_trash`` behaviour with linked accounts on a real
	site was not checked; the order below (transactions, settings, tax templates, accounts
	leaves-first, warehouses, cost centres, company) is the safe one and each step is
	wrapped so a refusal leaves a readable message instead of a half-deleted site.
	"""
	frappe = _frappe()
	if frappe is None:
		return []
	problems: list[str] = []

	def delete_all(doctype: str, filters: dict[str, Any], order_by: str | None = None) -> None:
		try:
			names = frappe.get_all(doctype, filters=filters, pluck="name", order_by=order_by)
		except Exception as exc:  # noqa: BLE001 - DocType may not exist on this site
			problems.append(f"{doctype}: {exc}")
			return
		for row in names:
			try:
				frappe.delete_doc(doctype, row, force=True, ignore_permissions=True)
			except Exception as exc:  # noqa: BLE001 - reported, see docstring
				problems.append(f"{doctype} {row}: {exc}")

	def delete_tree(doctype: str, parent_field: str) -> None:
		"""Leaves first, then groups whose children are gone; nested sets are not relied on."""
		try:
			remaining = set(frappe.get_all(doctype, filters={"company": name}, pluck="name"))
		except Exception as exc:  # noqa: BLE001
			problems.append(f"{doctype}: {exc}")
			return
		progress = True
		while remaining and progress:
			progress = False
			for row in sorted(remaining):
				if frappe.db.exists(doctype, {parent_field: row, "company": name}):
					continue
				try:
					frappe.delete_doc(doctype, row, force=True, ignore_permissions=True)
				except Exception as exc:  # noqa: BLE001 - reported, see docstring
					problems.append(f"{doctype} {row}: {exc}")
				remaining.discard(row)
				progress = True
		for row in sorted(remaining):
			problems.append(f"{doctype} {row}: could not be deleted (children remain)")

	# The simulation's Nyabo Documents are placeholders, not received primary documents: the
	# retention guard (art. 11.1) is lifted for this cleanup with the same flag maintenance
	# jobs use (``compliance.hooks.block_retained_document_delete``).
	frappe.flags.nyabo_allow_file_delete = True
	try:
		for doctype in (
			"Nyabo LLM Call",
			"Nyabo Proposal",
			"Nyabo Document",
			"Nyabo Rule",
			"Nyabo Account Alias",
			"Nyabo Company Settings",
		):
			delete_all(doctype, {"company": name})
	finally:
		frappe.flags.pop("nyabo_allow_file_delete", None)
	for doctype in (
		"Sales Taxes and Charges Template",
		"Purchase Taxes and Charges Template",
		"Item Tax Template",
	):
		delete_all(doctype, {"company": name})
	delete_tree("Account", "parent_account")
	delete_tree("Warehouse", "parent_warehouse")
	delete_tree("Cost Center", "parent_cost_center")
	try:
		if frappe.db.exists("Company", name):
			frappe.delete_doc("Company", name, force=True, ignore_permissions=True)
	except Exception as exc:  # noqa: BLE001
		problems.append(f"Company {name}: {exc}")
	for problem in problems:
		print(mn.SIM_COMPANY_CLEANUP_FAILED.format(company=name, error=problem))
	return problems


def _record_on_site(
	company: str, receipt: dict[str, Any], outcome: harness.ProposalOutcome, use_site: bool | None = None
) -> str | None:
	"""A Nyabo Document + Proposal pair so the site shows what was proposed; None off-site."""
	frappe = _frappe(use_site)
	if frappe is None:
		return None
	document = frappe.get_doc(
		{
			"doctype": "Nyabo Document",
			"company": company,
			"doc_type": "receipt",
			"status": "proposed",
			"file_hash": f"sim:{company}:{receipt.get('receipt_id') or receipt.get('seller_name')}",
			"file": "/private/files/simulation-placeholder.jpg",
			"mime_type": "image/jpeg",
		}
	)
	document.insert(ignore_permissions=True)
	proposal = frappe.get_doc(
		{
			"doctype": "Nyabo Proposal",
			"document": document.name,
			"company": company,
			"kind": "receipt",
			"status": "proposed",
			"needs_accountant": 1 if outcome.needs_accountant else 0,
			"posting_date": outcome.entry.posting_date.isoformat() if outcome.entry else None,
			"total": float(outcome.entry.total) if outcome.entry else 0,
			"vat_amount": float(outcome.entry.vat_amount) if outcome.entry else 0,
			"vat_treatment": outcome.vat_treatment,
			"account_code": outcome.account_code,
			"explanation": outcome.explanation[:200],
			"citation": _citation_text(outcome),
			"entry_json": json.dumps(outcome.entry.to_dict(), ensure_ascii=False) if outcome.entry else None,
			"extracted_json": json.dumps(receipt, ensure_ascii=False, default=str),
			"warnings_json": json.dumps(list(outcome.warnings), ensure_ascii=False),
			"model": outcome.llm[-1].model if outcome.llm else "mock-model",
		}
	)
	proposal.insert(ignore_permissions=True)
	return proposal.name


# --- rendering ------------------------------------------------------------------------------------


def _citation_text(outcome: harness.ProposalOutcome) -> str:
	if outcome.citation is None:
		return ""
	section = outcome.citation.section or mn.CITATION_SECTION_PENDING
	return f"{outcome.citation.instrument}, {section}"


def render_column(
	regime: str, on_date: dt.date, outcome: harness.ProposalOutcome, leaves: dict[str, str]
) -> list[str]:
	lines = [mn.SIM_COLUMN_HEADER.format(regime=REGIME_LABELS.get(regime, regime), date=on_date.isoformat())]
	if outcome.entry is None:
		lines.append(mn.SIM_NO_ENTRY.format(reason="; ".join(outcome.problems) or "-"))
	else:
		lines.append(mn.SIM_DOCUMENT_KIND.get(outcome.entry.document_kind, outcome.entry.document_kind))
		for line in outcome.entry.lines:
			side = mn.DEBIT_SHORT if line.debit > 0 else mn.CREDIT_SHORT
			amount = line.debit if line.debit > 0 else line.credit
			lines.append(
				mn.SIM_LINE.format(
					side=side,
					code=line.account_code,
					name=leaves.get(line.account_code, ""),
					amount=fmt_mnt(amount),
				)
			)
	lines.append(
		mn.SIM_VAT_LINE.format(
			treatment=mn.VAT_TREATMENT_LABELS.get(outcome.vat_treatment, outcome.vat_treatment)
		)
	)
	lines.append(mn.SIM_EXPLANATION.format(explanation=outcome.reason_mn))
	lines.append(mn.SIM_CITATION.format(citation=_citation_text(outcome) or mn.CITATION_SECTION_PENDING))
	lines.append(mn.SIM_NEEDS_ACCOUNTANT if outcome.needs_accountant else mn.SIM_AUTO_OK)
	for flag in outcome.flags:
		key = flag.split(":", 1)[0]
		label = mn.EVAL_FLAG_LABELS.get(key, key)
		if ":" in flag:
			label = f"{label} ({mn.FIELD_LABELS.get(flag.split(':', 1)[1], flag.split(':', 1)[1])})"
		lines.append(mn.CARD_WARNING_LINE.format(warning=label))
	return lines


def _wrap(text: str, width: int) -> list[str]:
	words = text.split()
	out: list[str] = []
	current = ""
	for word in words:
		candidate = f"{current} {word}".strip()
		if len(candidate) > width and current:
			out.append(current)
			current = word
		else:
			current = candidate
	if current or not out:
		out.append(current)
	return out


def side_by_side(columns: Sequence[list[str]], width: int = COLUMN_WIDTH) -> str:
	wrapped = [[w for line in col for w in _wrap(line, width)] for col in columns]
	height = max(len(c) for c in wrapped)
	rows = []
	for i in range(height):
		cells = [(col[i] if i < len(col) else "").ljust(width) for col in wrapped]
		rows.append(" │ ".join(cells).rstrip())
	return "\n".join(rows)


# --- entry points --------------------------------------------------------------------------------------


def simulate_receipt(
	case_or_path: str | Path | EvalCase | dict[str, Any],
	regimes: Iterable[str] = DEFAULT_REGIMES,
	dates: Iterable[str] = DEFAULT_DATES,
	*,
	client: MockLlmClient | None = None,
	print_output: bool = True,
	keep_companies: bool = False,
	use_site: bool | None = None,
) -> str:
	"""Print (and return) the two proposed entries side by side with citations.

	``use_site=None`` uses a connected site when there is one (throwaway ``SIM-`` companies,
	Nyabo Document / Proposal rows), ``False`` runs pure Python, ``True`` insists on a site.
	"""
	frappe = _frappe(use_site)
	if use_site and frappe is None:
		raise RuntimeError("simulate_receipt(use_site=True) needs a connected Frappe site")
	if frappe is not None:
		frappe.flags.nyabo_simulation = True
	client = client or MockLlmClient()
	case = load_case(case_or_path)
	receipt, input_json = _receipt_from_case(case, client)
	seller_vat_payer, lookup_source = seller_lookup(receipt)
	if input_json.get("seller_vat_payer") is not None:
		seller_vat_payer = input_json["seller_vat_payer"]
	rules = harness.load_rules(str(input_json.get("scheme") or harness.DEFAULT_SCHEME))
	columns: list[list[str]] = []
	companies: list[str] = []
	try:
		for regime, date_text in zip(regimes, dates, strict=True):
			on_date = dt.date.fromisoformat(date_text)
			company = provision_sim_company(regime, use_site) or f"{SIM_PREFIX}{regime}"
			if company.startswith(SIM_PREFIX) and frappe is not None:
				companies.append(company)
			per_regime = dict(input_json)
			by_regime = input_json.get("classify_fixture_by_regime") or {}
			if regime in by_regime:
				per_regime["classify_fixture"] = by_regime[regime]
			per_regime["receipt"] = dict(receipt, date=on_date.isoformat())
			per_regime["seller_vat_payer"] = seller_vat_payer
			spec = harness.ProposeInput.from_case(per_regime, regime, on_date, company)
			outcome = harness.propose(spec, client)
			_record_on_site(company, per_regime["receipt"], outcome, use_site)
			columns.append(render_column(regime, on_date, outcome, dict(rules.leaves)))
	finally:
		if not keep_companies:
			for company in companies:
				cleanup_sim_company(company)
	title = mn.SIM_TITLE.format(case=case.case_id)
	subtitle = f"{receipt.get('seller_name')} · {fmt_mnt(receipt.get('total') or 0)}{'₮' if input_json.get('currency', 'MNT') == 'MNT' else ' ' + str(input_json.get('currency'))} · seller lookup: {lookup_source}"
	text = "\n".join([title, subtitle, "", side_by_side(columns)])
	if print_output:
		print(text)
	return text


def simulate_statement(path: str | Path, *, print_output: bool = True) -> str:
	"""Parse a statement export with the seed layouts and show what each line would become.

	Uses ``nyabo_mn.parsers.excel.read_rows`` when that module exists (owned by another
	agent); otherwise reads CSV directly. Layout detection and matching are the core
	modules; with no vouchers to match against, lines come out as fee / transfer / none.
	"""
	from nyabo_mn.core import matching, statements
	from nyabo_mn.nyabo.seed import load_seed

	path = Path(path)
	rows = _read_rows(path)
	layouts = [statements.LayoutSpec.from_dict(row) for row in load_seed("bank_layouts")["rows"]]
	layout = statements.detect_layout(rows, layouts)
	if layout is None:
		layout = statements.guess_layout(rows)
	lines_out = [mn.SIM_STATEMENT_TITLE.format(path=path.name)]
	if layout is None:
		lines_out.append(mn.SIM_STATEMENT_NO_LAYOUT)
		text = "\n".join(lines_out)
		if print_output:
			print(text)
		return text
	bank_lines = statements.parse_rows(rows, layout)
	pairs = matching.pair_transfers(bank_lines)
	paired = {line.row_hash for a, b in pairs for line in (a, b)}
	for line in bank_lines:
		if line.row_hash in paired:
			result = mn.MATCH_REASON_TRANSFER
		else:
			result = matching.pick(line, []).reason
		lines_out.append(
			mn.SIM_STATEMENT_LINE.format(
				date=line.date.isoformat(),
				amount=fmt_mnt(line.amount),
				description=line.description[:40],
				result=result,
			)
		)
	text = "\n".join(lines_out)
	if print_output:
		print(text)
	return text


def _read_rows(path: Path) -> list[list[Any]]:
	"""``parsers.excel.read_rows(data, filename)`` — the bot's own reader — else a plain csv/xlsx read."""
	try:
		excel = importlib.import_module("nyabo_mn.parsers.excel")
	except ImportError:
		excel = None
	if excel is not None:
		return list(excel.read_rows(path.read_bytes(), path.name))
	if path.suffix.lower() == ".csv":
		import csv

		with path.open(encoding="utf-8-sig", newline="") as fh:
			return [list(row) for row in csv.reader(fh)]
	import openpyxl

	book = openpyxl.load_workbook(path, read_only=True, data_only=True)
	sheet = book[book.sheetnames[0]]
	return [list(row) for row in sheet.iter_rows(values_only=True)]


def run(case: str = "petrovis_fuel", statement: str | None = None) -> str:
	"""bench entry point: ``--kwargs '{"case": "petrovis_fuel"}'`` or ``{"statement": "<path>"}``."""
	if statement:
		return simulate_statement(statement)
	return simulate_receipt(case)


def main(argv: Sequence[str] | None = None) -> int:
	parser = argparse.ArgumentParser(description="Simulate a receipt under both regimes")
	parser.add_argument("case", nargs="?", default="petrovis_fuel", help="golden case id or JSON path")
	parser.add_argument("--statement", default=None, help="xlsx/csv statement to simulate instead")
	args = parser.parse_args(argv)
	run(args.case, args.statement)
	return 0


if __name__ == "__main__":
	sys.exit(main())


__all__ = [
	"DEFAULT_DATES",
	"DEFAULT_REGIMES",
	"SIM_PREFIX",
	"cleanup_sim_company",
	"load_case",
	"provision_sim_company",
	"render_column",
	"run",
	"side_by_side",
	"simulate_receipt",
	"simulate_statement",
]
