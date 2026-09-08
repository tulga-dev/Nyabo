"""VAT summary, simplified 1% summary, month-end checklist / trial balance / summaries."""

from __future__ import annotations

from decimal import Decimal

import frappe
import pytest
from compliance_helpers import BANK, CASH, EXPENSE, INCOME, INPUT_VAT, OUTPUT_VAT, make_je, make_pi, make_si
from openpyxl import load_workbook

from nyabo_mn.core.money import fmt_mnt
from nyabo_mn.i18n import mn
from nyabo_mn.reports import export, month_end, rules_bridge, simplified_summary, vat_summary


@pytest.fixture
def books(company):
	"""A VAT-registered March: one purchase invoice with input VAT, one sales invoice with output VAT."""
	supplier = frappe.get_doc({"doctype": "Supplier", "supplier_name": "Петровис ХХК"}).insert()
	customer = frappe.get_doc({"doctype": "Customer", "customer_name": "Хэрэглэгч ХХК"}).insert()
	pi = make_pi(
		company,
		supplier,
		nyabo_primary_document_ref="AB-1",
		taxes=[
			{"account_head": INPUT_VAT, "charge_type": "On Net Total", "rate": 10, "description": "НӨАТ 10%"}
		],
	).insert()
	pi.submit()
	si = make_si(company, customer, amount=200000, nyabo_primary_document_ref="SI-1").insert()
	si.submit()
	je = make_je(
		company,
		amount=50000,
		posting_date="2026-03-20",
		debit=EXPENSE,
		credit=BANK,
		nyabo_primary_document_ref="x",
	)
	je.insert()
	je.submit()
	return company


def _settings(company, regime):
	"""Replace the regime history (provisioning wrote the first row) with one open row, or none."""
	name = frappe.db.get_value("Nyabo Company Settings", {"company": company}, "name")
	doc = (
		frappe.get_doc("Nyabo Company Settings", name)
		if name
		else frappe.get_doc({"doctype": "Nyabo Company Settings", "company": company})
	)
	doc.regimes = []
	if regime:
		doc.append("regimes", {"regime": regime, "effective_from": "2026-01-01"})
	doc.accountant_of_record_name = "Б. Батаа"
	doc.save()
	return doc


def _seed_rows(key: str):
	"""The shipped seed rows of one key as Nyabo Tax Parameter rows.

	The first row in the table makes it win over the seed file for *every* key, so a test that
	inserts its own row must also supply the other keys the summary reads (the regime's
	eligibility conditions).
	"""
	from nyabo_mn.nyabo.seed import load_seed
	from nyabo_mn.rules.seed import tax_parameter_values

	return [
		frappe.get_doc({"doctype": "Nyabo Tax Parameter", **tax_parameter_values(row)}).insert()
		for row in load_seed("tax_parameters")["rows"]
		if row["key"] == key
	]


def _rate_row(verified: int):
	"""An explicit simplified.rate row: the shipped seed row is verified, the guard test needs both."""
	for key in simplified_summary.ELIGIBILITY_KEYS:
		if not frappe.db.exists("Nyabo Tax Parameter", {"key": key}):
			_seed_rows(key)
	return frappe.get_doc(
		{
			"doctype": "Nyabo Tax Parameter",
			"key": "simplified.rate",
			"effective_from": "2026-01-01",
			"value_json": "0.01",
			"unit": "fraction",
			"status": "active",
			"verified": verified,
			"source_text": "ААНОАТ-ын тухай хууль",
		}
	).insert()


def test_vat_summary_from_gl_rows(books):
	summary = vat_summary.compute(books, "2026-03")
	assert summary["output_vat"] == Decimal("20000.00")
	assert summary["input_vat"] == Decimal("8500.00")
	assert summary["net"] == Decimal("11500.00")
	assert summary["accounts"] == {"output_vat": OUTPUT_VAT, "input_vat": INPUT_VAT}
	docs = {d["voucher_type"]: d for d in summary["by_document"]}
	assert docs["Sales Invoice"]["output_vat"] == Decimal("20000.00")
	assert docs["Purchase Invoice"]["input_vat"] == Decimal("8500.00")
	assert vat_summary.compute(books, "2026-04")["net"] == Decimal("0.00")


def test_simplified_summary_uses_the_parameter_row_and_refuses_unverified(books):
	# the site's own (unverified) row wins over the shipped seed row, and is refused for a real figure
	row = _rate_row(verified=0)
	with pytest.raises(rules_bridge.UnverifiedRuleError) as exc:
		simplified_summary.compute(books, "2026-Q1")
	assert "simplified.rate" in str(exc.value)
	assert isinstance(exc.value, frappe.ValidationError)
	simulated = simplified_summary.compute(books, "2026-Q1", simulation=True)
	assert simulated["revenue"] == Decimal("200000.00")
	assert simulated["tax_1pct"] == Decimal("2000.00")
	assert simulated["simulation"] is True and simulated["rate_row"]["verified"] is False
	assert [m["period"] for m in simulated["months"]] == ["2026-01", "2026-02", "2026-03"]
	assert simulated["months"][2]["revenue"] == Decimal("200000.00")

	row.verified = 1  # the admin checked the primary text
	row.save()
	real = simplified_summary.compute(books, "2026-Q1")
	assert real["tax_1pct"] == Decimal("2000.00") and real["rate_row"]["verified"] is True
	with pytest.raises(frappe.ValidationError):
		simplified_summary.compute(books, "2026-Q9")


def test_simplified_summary_reads_the_shipped_seed_before_the_first_sync(books):
	"""No Nyabo Tax Parameter row yet: the seed file's (verified) simplified.rate drives the figure."""
	real = simplified_summary.compute(books, "2026-Q1")
	assert real["tax_1pct"] == Decimal("2000.00") and real["rate_row"]["verified"] is True
	assert real["rate_row"]["article"]


def test_month_end_checklist_and_trial_balance(books):
	frappe.get_doc(
		{"doctype": "Nyabo Proposal", "company": books, "kind": "receipt", "status": "proposed"}
	).insert()
	frappe.get_doc(
		{"doctype": "Supplier", "supplier_name": "Шинэ ХХК", "nyabo_pending_confirmation": 1}
	).insert()
	check = month_end.checklist(books, "2026-03")
	assert (check["proposals"], check["pending_suppliers"], check["unverified_rules"]) == (1, 1, 0)
	assert check["blocking"] is False and check["inventory_count_required"] is False
	assert "Шийдвэрлээгүй санал: 1" in check["text"]
	assert mn.MSG_CLOSE_INVENTORY_COUNT in month_end.checklist(books, "2026-12")["text"]

	tb = month_end.trial_balance(books, "2026-03")
	assert tb["source"] == "gl_entry" and tb["debit"] == tb["credit"] == Decimal("363500.00")
	by_account = {r["account"]: r for r in tb["rows"]}
	assert (
		by_account[EXPENSE]["debit"] == Decimal("135000.00") and by_account[EXPENSE]["account_code"] == "6210"
	)
	assert by_account[INCOME]["credit"] == Decimal("200000.00")
	assert by_account[CASH]["credit"] == Decimal("0.00") if CASH in by_account else True


def test_month_end_summaries_pick_the_regime(books):
	# provisioning recorded simplified 1% (vat_registered=0); without any regime row the close says so
	assert month_end.summaries(books, "2026-03", simulation=True)["is_vat_payer"] is False
	_settings(books, None)
	out = month_end.summaries(books, "2026-03")
	assert out["is_vat_payer"] is None and out["pdfs"] == {}
	assert any("2026-03-31" in line for line in out["text_lines"])
	_settings(books, "vat_payer")
	out = month_end.summaries(books, "2026-03")
	assert out["is_vat_payer"] is True and set(out["pdfs"]) == {"vat_summary"}
	assert any(fmt_mnt(11500) in line for line in out["text_lines"])
	assert mn.REPORT_VAT_SUMMARY in out["pdfs"]["vat_summary"].decode("utf-8")
	assert "DejaVu Sans" in out["pdfs"]["vat_summary"].decode("utf-8")

	_settings(books, "simplified_1pct")
	out = month_end.summaries(books, "2026-03", simulation=True)
	assert out["is_vat_payer"] is False and set(out["pdfs"]) == {"simplified_summary"}
	assert any("1% татвар" in line for line in out["text_lines"])
	html = out["pdfs"]["simplified_summary"].decode("utf-8")
	assert mn.LBL_SIMULATION in html and mn.JOURNAL_KEPT_BY in html
	_rate_row(verified=0)  # the site's own unverified rate row: refused outside the Simulation filter
	with pytest.raises(rules_bridge.UnverifiedRuleError):
		month_end.summaries(books, "2026-03")


def test_report_to_xlsx_round_trips(books):
	data = export.report_to_xlsx(
		"Nyabo VAT Summary", {"company": books, "from_date": "2026-03-01", "to_date": "2026-03-31"}
	)
	from io import BytesIO

	sheet = load_workbook(BytesIO(data)).active
	rows = [list(r) for r in sheet.iter_rows(values_only=True)]
	assert rows[0][:2] == [mn.COL_DATE, mn.COL_VOUCHER_TYPE]
	assert len(rows) == 3 and {r[1] for r in rows[1:]} == {"Purchase Invoice", "Sales Invoice"}


def test_trial_balance_falls_back_when_the_erpnext_report_refuses(books, monkeypatch):
	"""F7: ERPNext raises PermissionError / ValidationError, not DoesNotExistError — still a GL fallback.

	The Telegram user holds only ``Nyabo Accountant``, so ``get_report_doc("Trial Balance")``
	throws ``frappe.PermissionError``; a site whose Fiscal Year does not reach the period
	raises ``FiscalYearError`` (a ``ValidationError``). Neither may break ``/хаалт``.
	"""
	from erpnext.accounts import utils as erpnext_utils
	from frappe.desk import query_report

	def _refuse(*args, **kwargs):
		raise frappe.PermissionError("You don't have access to Report: Trial Balance")

	monkeypatch.setattr(query_report, "run", _refuse)
	tb = month_end.trial_balance(books, "2026-03")
	assert tb["source"] == "gl_entry" and tb["debit"] == tb["credit"] == Decimal("363500.00")
	assert tb["note"] == mn.MSG_CLOSE_TRIAL_BALANCE_SOURCE_FALLBACK

	def _no_fiscal_year(*args, **kwargs):
		raise erpnext_utils.FiscalYearError("2026-03-31 is not in any active Fiscal Year")

	monkeypatch.setattr(erpnext_utils, "get_fiscal_year", _no_fiscal_year)
	fallback = month_end.trial_balance(books, "2026-03")
	assert fallback["source"] == "gl_entry" and fallback["rows"]


def test_simplified_summary_resolves_the_regime_conditions_before_any_figure(books):
	"""F-03: art. 29.1 and 29.3.1 are looked up too, so a pending year refuses instead of computing."""
	summary = simplified_summary.compute(books, "2026-Q1")
	conditions = summary["eligibility"]["rows"]
	assert set(conditions) == set(simplified_summary.ELIGIBILITY_KEYS)
	assert conditions["simplified.revenue_threshold"]["article"] == "29.1"
	assert conditions["simplified.requires_not_vat_registered"]["article"] == "29.3.1"
	assert all(c["verified"] for c in conditions.values())
	assert summary["needs_accountant"] is False and summary["warnings"] == []
	assert summary["eligibility"]["threshold"] == Decimal("50000000")

	# the 2027 eligibility rows are pending on purpose: no 1% figure, in a simulation either
	for simulation in (False, True):
		with pytest.raises(frappe.ValidationError) as exc:
			simplified_summary.compute(books, "2027-Q1", simulation=simulation)
		assert "simplified." in str(exc.value)
	with pytest.raises(frappe.ValidationError):
		month_end.summaries(books, "2027-03")

	# a VAT withholding payer is outside the regime (29.3.1), whatever the revenue is
	_settings(books, "vat_payer")
	with pytest.raises(frappe.ValidationError) as exc:
		simplified_summary.compute(books, "2026-Q1")
	assert mn.MSG_SIMPLIFIED_NOT_ELIGIBLE_VAT in str(exc.value)


def test_simplified_summary_flags_a_prior_year_above_the_threshold(books):
	"""The books can raise the art. 29.1 question (they are not the confirmed return, so never settle it)."""
	if not frappe.db.exists("Fiscal Year", "2025"):
		frappe.get_doc(
			{
				"doctype": "Fiscal Year",
				"year": "2025",
				"year_start_date": "2025-01-01",
				"year_end_date": "2025-12-31",
			}
		).insert(ignore_permissions=True)
	si = make_si(
		books,
		"Хэрэглэгч ХХК",
		amount=60000000,
		posting_date="2025-06-15",
		vat=False,
		nyabo_primary_document_ref="SI-2025",
	).insert()
	si.submit()

	summary = simplified_summary.compute(books, "2026-Q1")
	assert summary["eligibility"]["prior_year_revenue"] == Decimal("60000000.00")
	assert summary["eligibility"]["over_threshold"] is True and summary["needs_accountant"] is True
	assert mn.WARN_SIMPLIFIED_OVER_THRESHOLD.split("{")[0] in summary["warnings"][0]

	out = month_end.summaries(books, "2026-03")
	assert out["is_vat_payer"] is False and summary["warnings"][0] in out["text_lines"]
	html = out["pdfs"]["simplified_summary"].decode("utf-8")
	assert mn.LBL_REGIME_CONDITION in html and "29.1" in html
